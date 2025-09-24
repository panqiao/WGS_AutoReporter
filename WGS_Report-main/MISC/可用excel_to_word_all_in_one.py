#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel to Word转换工具 - 完整版
功能：
1. 预处理Excel数据，拆分包含多个phenotypes的行
2. 按Phenotype和Drug组织数据
3. 生成格式化的Word文档报告
"""

import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from collections import defaultdict

# ==================== 预处理功能模块 ====================

def preprocess_phenotypes(input_file, output_file):
    """
    预处理Excel文件，将包含多个phenotypes的行拆分成独立的行
    
    参数:
    input_file: 原始Excel文件路径
    output_file: 处理后的Excel文件路径
    """
    
    print("读取原始Excel文件...")
    df = pd.read_excel(input_file)
    print(f"原始数据: {len(df)} 行")
    
    # 创建新的数据列表
    new_rows = []
    
    # 统计信息
    split_count = 0
    multi_phenotype_rows = 0
    
    # 遍历每一行
    for idx, row in df.iterrows():
        phenotype_value = row['Phenotype(s)']
        
        # 检查是否为空值
        if pd.isna(phenotype_value):
            # 保留空值行
            new_rows.append(row.to_dict())
            continue
        
        # 转换为字符串
        phenotype_str = str(phenotype_value)
        
        # 检查是否包含分号
        if ';' in phenotype_str:
            multi_phenotype_rows += 1
            # 拆分phenotypes
            phenotypes = [p.strip() for p in phenotype_str.split(';')]
            split_count += len(phenotypes) - 1  # 减1是因为原始行算1个
            
            # 为每个phenotype创建新行
            for phenotype in phenotypes:
                new_row = row.to_dict()  # 复制整行数据
                new_row['Phenotype(s)'] = phenotype  # 更新Phenotype列
                new_rows.append(new_row)
        else:
            # 不包含分号的行直接添加
            new_rows.append(row.to_dict())
    
    # 创建新的DataFrame
    new_df = pd.DataFrame(new_rows)
    
    print(f"\n处理统计:")
    print(f"  - 包含多个phenotypes的行: {multi_phenotype_rows}")
    print(f"  - 拆分产生的新行: {split_count}")
    print(f"  - 处理后总行数: {len(new_df)} 行")
    
    # 保存到新的Excel文件
    new_df.to_excel(output_file, index=False)
    print(f"\n预处理后的文件已保存至: {output_file}")
    
    # 显示拆分示例
    print("\n拆分示例:")
    example_count = 0
    for idx, row in df.iterrows():
        if pd.notna(row['Phenotype(s)']) and ';' in str(row['Phenotype(s)']):
            print(f"\n原始 Phenotype: {row['Phenotype(s)']}")
            print(f"对应药物: {row['Drug(s)']}")
            phenotypes = [p.strip() for p in str(row['Phenotype(s)']).split(';')]
            print(f"拆分后: {phenotypes}")
            example_count += 1
            if example_count >= 3:  # 只显示3个示例
                break
    
    return new_df

def validate_preprocessing(original_file, processed_file):
    """
    验证预处理的正确性
    """
    print("\n" + "="*60)
    print("验证预处理结果...")
    
    original_df = pd.read_excel(original_file)
    processed_df = pd.read_excel(processed_file)
    
    # 检查列是否一致
    if list(original_df.columns) == list(processed_df.columns):
        print("✓ 列名保持一致")
    else:
        print("✗ 列名不一致")
        return False
    
    # 检查非Phenotype列的数据是否正确复制
    test_row = original_df[original_df['Phenotype(s)'].str.contains(';', na=False)].iloc[0] if any(original_df['Phenotype(s)'].str.contains(';', na=False)) else None
    
    if test_row is not None:
        original_phenotype = test_row['Phenotype(s)']
        phenotypes = [p.strip() for p in original_phenotype.split(';')]
        
        # 检查拆分后的行
        for phenotype in phenotypes:
            matched_rows = processed_df[
                (processed_df['Phenotype(s)'] == phenotype) & 
                (processed_df['Drug(s)'] == test_row['Drug(s)']) &
                (processed_df['Gene'] == test_row['Gene'])
            ]
            
            if len(matched_rows) > 0:
                print(f"✓ Phenotype '{phenotype}' 正确拆分并保留了原始数据")
            else:
                print(f"✗ Phenotype '{phenotype}' 拆分可能有问题")
    
    print("="*60)
    return True

# ==================== Word文档生成模块 ====================

def read_excel_file(file_path):
    """读取Excel文件"""
    df = pd.read_excel(file_path)
    return df

def organize_data_by_phenotype_and_drug(df):
    """
    按Phenotype和Drug组织数据
    返回嵌套字典结构：{phenotype: {drug: [rows]}}
    """
    organized_data = defaultdict(lambda: defaultdict(list))
    
    for _, row in df.iterrows():
        phenotype = row['Phenotype(s)']
        drug = row['Drug(s)']
        
        # 将每一行的相关信息存储
        row_data = {
            'Result': row['Result'],
            'Gene': row['Gene'],
            'Variant/Haplotypes': row['Variant/Haplotypes'],
            'Genotype/Allele': row['Genotype/Allele'],
            'Annotation Text': row['Annotation Text']
        }
        
        organized_data[phenotype][drug].append(row_data)
    
    return organized_data

def create_word_document(organized_data, output_path):
    """创建Word文档并写入组织好的数据"""
    doc = Document()
    
    # 设置文档标题
    title = doc.add_heading('Pharmacogenomics Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()  # 添加空行
    
    # 遍历每个Phenotype（处理NaN值）
    phenotypes = list(organized_data.keys())
    # 将NaN转换为字符串"Unknown"以便排序
    phenotypes_sorted = sorted(phenotypes, key=lambda x: str(x) if pd.notna(x) else "Unknown")
    
    for phenotype_idx, phenotype in enumerate(phenotypes_sorted, 1):
        # 处理NaN值显示
        phenotype_display = str(phenotype) if pd.notna(phenotype) else "Unknown"
        # 添加Phenotype作为一级标题
        phenotype_heading = doc.add_heading(f"{phenotype_idx}. {phenotype_display}", level=1)
        
        # 遍历该Phenotype下的每个Drug
        drug_data = organized_data[phenotype]
        drugs = list(drug_data.keys())
        drugs_sorted = sorted(drugs, key=lambda x: str(x) if pd.notna(x) else "Unknown")
        
        for drug_idx, drug in enumerate(drugs_sorted, 1):
            # 处理NaN值显示
            drug_display = str(drug) if pd.notna(drug) else "Unknown"
            # 将药物名称中的分号替换为&符号，更美观
            drug_display = drug_display.replace(';', ' & ')
            # 添加Drug作为二级标题（药物名称只显示一次）
            drug_heading = doc.add_heading(f"{phenotype_idx}.{drug_idx} {drug_display}", level=2)
            
            # 获取该药物的所有记录
            records = drug_data[drug]
            
            # 如果有多条记录，显示记录数量
            if len(records) > 1:
                count_para = doc.add_paragraph()
                count_para.add_run(f"(共 {len(records)} 条记录)").italic = True
            
            # 遍历每条记录
            for record_idx, record in enumerate(records, 1):
                # 如果有多条记录，添加记录编号
                if len(records) > 1:
                    record_para = doc.add_paragraph()
                    record_run = record_para.add_run(f"记录 {record_idx}:")
                    record_run.bold = True
                    record_run.font.color.rgb = RGBColor(0, 0, 139)  # 深蓝色
                
                # 添加Result
                result_para = doc.add_paragraph()
                result_para.add_run("Result: ").bold = True
                result_text = str(record['Result']) if pd.notna(record['Result']) else "N/A"
                result_para.add_run(result_text)
                
                # 创建表格显示Gene、Variant/Haplotypes、Genotype/Allele
                table = doc.add_table(rows=2, cols=3)
                table.style = 'Light Grid'
                
                # 设置表头
                header_cells = table.rows[0].cells
                headers = ['Gene', 'Variant/Haplotypes', 'Genotype/Allele']
                for i, header in enumerate(headers):
                    header_cells[i].text = header
                    # 设置表头样式
                    for paragraph in header_cells[i].paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                
                # 填充数据行
                data_cells = table.rows[1].cells
                data_cells[0].text = str(record['Gene']) if pd.notna(record['Gene']) else "N/A"
                data_cells[1].text = str(record['Variant/Haplotypes']) if pd.notna(record['Variant/Haplotypes']) else "N/A"
                data_cells[2].text = str(record['Genotype/Allele']) if pd.notna(record['Genotype/Allele']) else "N/A"
                
                # 添加Annotation Text
                annotation_para = doc.add_paragraph()
                annotation_para.add_run("Annotation Text: ").bold = True
                annotation_text = str(record['Annotation Text']) if pd.notna(record['Annotation Text']) else "N/A"
                
                # 如果文本过长，进行换行处理
                if len(annotation_text) > 200:
                    # 分段显示长文本
                    annotation_para.add_run(annotation_text[:200])
                    remaining_text = annotation_text[200:]
                    while remaining_text:
                        doc.add_paragraph(remaining_text[:200], style='Normal')
                        remaining_text = remaining_text[200:]
                else:
                    annotation_para.add_run(annotation_text)
                
                # 在记录之间添加分隔线（如果不是最后一条记录）
                if record_idx < len(records):
                    separator = doc.add_paragraph("－" * 30)
                    separator.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # 在不同药物之间添加空行
            doc.add_paragraph()
    
    # 添加文档结尾信息
    doc.add_page_break()
    end_para = doc.add_paragraph()
    end_para.add_run("文档生成完成").bold = True
    end_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 保存文档
    doc.save(output_path)
    print(f"Word文档已保存至: {output_path}")
    return output_path

def generate_summary(organized_data):
    """生成数据摘要"""
    print("\n" + "="*60)
    print("数据摘要：")
    print("="*60)
    
    total_records = 0
    phenotype_count = len(organized_data)
    
    for phenotype in organized_data:
        drug_count = len(organized_data[phenotype])
        phenotype_records = sum(len(records) for records in organized_data[phenotype].values())
        total_records += phenotype_records
        print(f"\n{phenotype}:")
        print(f"  - 药物数量: {drug_count}")
        print(f"  - 记录总数: {phenotype_records}")
        
        # 显示每个药物的记录数
        for drug in sorted(organized_data[phenotype].keys()):
            record_count = len(organized_data[phenotype][drug])
            if record_count > 1:
                print(f"    • {drug}: {record_count} 条记录")
    
    print("\n" + "="*60)
    print(f"总计: {phenotype_count} 个Phenotypes, {total_records} 条记录")
    print("="*60)

# ==================== 主程序 ====================

def main():
    # 文件路径
    original_excel = "/Users/qiaopan/Downloads/leo-matched_snp_genotypes_results666_reordered 2.xlsx"
    preprocessed_excel = "/Users/qiaopan/Library/Mobile Documents/com~apple~CloudDocs/Work-Ebovir/全基因组测序项目/pharmacy/preprocessed_data.xlsx"
    output_file = "/Users/qiaopan/Library/Mobile Documents/com~apple~CloudDocs/Work-Ebovir/全基因组测序项目/pharmacy/pharmacogenomics_report.docx"
    
    try:
        # 步骤1: 预处理数据（拆分含有多个phenotypes的行）
        print("="*60)
        print("步骤1: 预处理数据")
        print("="*60)
        preprocessed_df = preprocess_phenotypes(original_excel, preprocessed_excel)
        print("\n预处理完成！")
        
        # 验证预处理结果
        validate_preprocessing(original_excel, preprocessed_excel)
        
        # 步骤2: 读取预处理后的Excel文件
        print("\n" + "="*60)
        print("步骤2: 生成Word文档")
        print("="*60)
        print("正在读取预处理后的Excel文件...")
        df = read_excel_file(preprocessed_excel)
        print(f"成功读取 {len(df)} 条记录")
        
        print("\n正在组织数据...")
        organized_data = organize_data_by_phenotype_and_drug(df)
        
        # 生成并显示数据摘要
        generate_summary(organized_data)
        
        print("\n正在创建Word文档...")
        output_path = create_word_document(organized_data, output_file)
        
        print("\n✅ 完成！")
        print(f"📄 Word文档路径: {output_path}")
        
    except Exception as e:
        print(f"\n❌ 发生错误: {str(e)}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()